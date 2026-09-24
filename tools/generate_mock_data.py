"""Generate the fictional demo dataset for this Power BI project.

All customers, salespeople, suppliers, brands, products, addresses, amounts and
dates are synthetic. The files keep the exact column layout, formats and sign
conventions of the original ERP exports so every Power Query step and DAX
measure in the model runs unchanged.

Usage:  python tools/generate_mock_data.py        (writes into ../Data)
"""
import csv
import json
import math
import random
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

SEED = 2026
AS_OF = date(2026, 9, 23)
START = date(2025, 1, 1)
PO_START = date(2024, 11, 1)
ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / "Data"
TAX = json.loads((ROOT / "taxonomy.json").read_text(encoding="utf-8"))
rng = random.Random(SEED)

COMPANIES = [
    dict(name="Crestline New Jersey Inc", wh="New Jersey Warehouse", inv="NJINV", rinv="RINV", so="NJS", po="NJP", share=0.36, customers=170, reps=12, in_sales_all=True),
    dict(name="Crestline Pacific Inc", wh="Los Angeles Warehouse", inv="CAINV", rinv="CARINV", so="CAS", po="CAP", share=0.34, customers=140, reps=8, in_sales_all=True),
    dict(name="Crestline Chicago Inc", wh="Chicago Warehouse", inv="ILINV", rinv="ILRINV", so="ILS", po="ILP", share=0.13, customers=55, reps=4, in_sales_all=False),
    dict(name="Crestline Houston Inc", wh="Houston-Warehouse", inv="TXINV", rinv="TXRINV", so="TXS", po="TXP", share=0.10, customers=45, reps=3, in_sales_all=False),
    dict(name="Crestline Northwest Inc", wh="Seattle Warehouse", inv="NWINV", rinv="NWRINV", so="NWS", po="NWP", share=0.07, customers=35, reps=3, in_sales_all=True),
]
EXTRA_WH = {"Crestline New Jersey Inc": "Georgia Warehouse", "Crestline Pacific Inc": "San Francisco Warehouse"}

STATE_ABBR = {
    "Alabama": "AL", "Arizona": "AZ", "Arkansas": "AR", "California": "CA", "Colorado": "CO", "Connecticut": "CT",
    "Delaware": "DE", "District of Columbia": "DC", "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS", "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME",
    "Maryland": "MD", "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS", "Missouri": "MO",
    "Montana": "MT", "Nebraska": "NE", "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM",
    "New York": "NY", "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR",
    "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC", "South Dakota": "SD", "Tennessee": "TN",
    "Texas": "TX", "Utah": "UT", "Vermont": "VT", "Virginia": "VA", "Washington": "WA", "West Virginia": "WV",
    "Wisconsin": "WI", "Wyoming": "WY",
}

FIRST = ["Kevin", "Lina", "Oscar", "Mia", "Ethan", "Grace", "Leo", "Ivy", "Ryan", "Nora", "Simon", "Cathy", "Victor",
         "Helen", "Jason", "Amy", "Derek", "Fiona", "Owen", "Tracy", "Bruce", "Joyce", "Frank", "Wendy", "Aaron",
         "Stella", "Colin", "Irene", "Jerry", "Rita", "Harvey", "Emma", "Louis", "Selina", "Peter", "Yvonne"]
SUR = ["Zhao", "Qian", "Sun", "Zhou", "Wu", "Zheng", "Feng", "Chu", "Wei", "Jiang", "Shen", "Han", "Yang", "Zhu",
       "Qin", "Xu", "He", "Lyu", "Shi", "Kong", "Cao", "Yan", "Hua", "Jin", "Tao", "Xie", "Zou", "Pan", "Ge", "Fan"]
GIVEN = ["Hao", "Ming", "Jie", "Lei", "Yun", "Ting", "Bo", "Xin", "Rui", "Qiang", "Yan", "Fang", "Kai", "Lan", "Peng"]

STORE_A = ["Golden", "Lucky", "Jade", "Pacific", "Harbor", "Evergreen", "Sunrise", "Bamboo", "Lotus", "Phoenix",
           "Silver", "Orchid", "Maple", "Crescent", "Summit", "Riverside", "Emerald", "Azure", "Coral", "Willow",
           "Pearl", "Morning", "Northstar", "Oak", "Cedar", "Blue Bay", "Red Maple", "Twin Pines", "Sakura", "Hibiscus"]
STORE_B = ["Fresh", "Asian", "Family", "Garden", "Village", "Harvest", "Oriental", "Seaside", "Valley", "Plaza",
           "Hometown", "Kitchen", "Grand", "Eastern", "Tokyo", "Seoul", "Saigon", "Canton", "Metro", "Corner"]
TYPE_SUFFIX = {
    "Supermarket/Grocery": ["Supermarket", "Market", "Grocery", "Food Market", "Mart"],
    "Wholesaler": ["Trading LLC", "Distribution Inc", "Wholesale Corp", "Foods Inc", "Supply Co"],
    "On-Premise": ["Restaurant", "Kitchen", "Noodle House", "Bistro", "Hot Pot"],
    "E-Commerce": ["Online Store", "eShop LLC", "Direct Inc"],
    "Others": ["Services LLC", "Group Inc", "Enterprises"],
    "Retail": ["Store", "Shop", "Convenience"],
    "Pharmacy": ["Pharmacy", "Drug Store"],
}
CUST_TYPES = [("Supermarket/Grocery", 50), ("Wholesaler", 12), ("On-Premise", 10), ("Others", 8), ("Retail", 6),
              ("E-Commerce", 5), ("Pharmacy", 3), ("Employee", 2)]
RACES = [("", 64), ("China", 16), ("Southeast Asian", 9), ("Mainstream", 6), ("Japan&Korea", 5)]
STREETS = ["MAIN ST", "BROADWAY", "OAK AVE", "PARK AVE", "ELM ST", "CENTER ST", "WASHINGTON AVE", "LINCOLN BLVD",
           "RIVER RD", "HILLSIDE AVE", "MARKET ST", "UNION TPKE", "SUNSET BLVD", "LAKE ST", "HIGHLAND AVE",
           "COMMERCE DR", "INDUSTRIAL WAY", "ROUTE 1", "GARDEN ST", "MAPLE AVE"]

SYL = [("金", "Jin"), ("福", "Fu"), ("源", "Yuan"), ("禾", "He"), ("山", "Shan"), ("岭", "Ling"), ("青", "Qing"),
       ("丰", "Feng"), ("嘉", "Jia"), ("乐", "Le"), ("晨", "Chen"), ("星", "Xing"), ("悦", "Yue"), ("味", "Wei"),
       ("鲜", "Xian"), ("香", "Xiang"), ("田", "Tian"), ("湾", "Wan"), ("川", "Chuan"), ("松", "Song"), ("竹", "Zhu"),
       ("梅", "Mei"), ("兰", "Lan"), ("安", "An"), ("宝", "Bao"), ("顺", "Shun"), ("瑞", "Rui"), ("润", "Run"),
       ("谷", "Gu"), ("园", "Yuan"), ("喜", "Xi"), ("友", "You"), ("亮", "Liang"), ("辰", "Chen"), ("溪", "Xi"),
       ("峰", "Feng"), ("云", "Yun"), ("鹿", "Lu"), ("桥", "Qiao"), ("舟", "Zhou"), ("橙", "Cheng"), ("麦", "Mai")]
FLAVORS = [("原味", "Original"), ("香辣", "Spicy"), ("麻辣", "Mala"), ("芒果", "Mango"), ("草莓", "Strawberry"),
           ("蜜桃", "Peach"), ("荔枝", "Lychee"), ("葡萄", "Grape"), ("柠檬", "Lemon"), ("抹茶", "Matcha"),
           ("牛奶", "Milk"), ("椰子", "Coconut"), ("蜂蜜", "Honey"), ("海盐", "Sea Salt"), ("烧烤", "BBQ"),
           ("番茄", "Tomato"), ("酸菜", "Pickled Vegetable"), ("牛肉", "Beef"), ("鸡肉", "Chicken"), ("海鲜", "Seafood"),
           ("蒜香", "Garlic"), ("五香", "Five Spice"), ("红枣", "Jujube"), ("绿茶", "Green Tea"), ("乌龙", "Oolong"),
           ("桂花", "Osmanthus"), ("菠萝", "Pineapple"), ("哈密瓜", "Melon"), ("香草", "Vanilla"),
           ("巧克力", "Chocolate"), ("红豆", "Red Bean"), ("玉米", "Corn"), ("芋头", "Taro"), ("黑糖", "Brown Sugar")]
CAT_NOUN = {"饮料": "Drink", "糖果零食": "Snack", "面条类干货": "Noodles", "调味品": "Seasoning",
            "冷冻和冷藏食品": "Frozen Food", "美容 & 健康 & 居家": "Care", "粮食作物": "Grain",
            "蔬菜类干货": "Dried Vegetables", "蜂蜜果酱": "Spread", "面粉&淀粉类": "Flour", "罐头干货": "Canned Food",
            "食用油": "Oil", "冲泡类": "Instant Drink", "汤&即食餐": "Instant Meal", "蛋": "Eggs"}
PRICE_RANGE = {"饮料": (12, 32), "糖果零食": (18, 48), "面条类干货": (16, 42), "调味品": (14, 60),
               "冷冻和冷藏食品": (38, 72), "美容 & 健康 & 居家": (3, 26), "粮食作物": (15, 60),
               "蔬菜类干货": (25, 120)}
# generated names that happen to match real-world brands
AVOID_BRANDS = {"YOUJIA"}
SUPPLIER_PLACES = ["Donghai", "Xinglin", "Baiyun", "Qinglan", "Hengshan", "Jiangbei", "Nanxi", "Yuhua", "Taoyuan",
                   "Lanting", "Fengcheng", "Songxi"]
SUPPLIER_WORDS = ["Orient", "Harvest", "Bridge", "Crown", "Rising Sun", "Silk Road", "Green Leaf", "Golden Grain"]


def pick(pairs):
    items, weights = zip(*pairs)
    return rng.choices(items, weights)[0]


def fmt_num(x):
    """Branch-sales style: integers without decimals, floats as Python repr."""
    if x is None:
        return ""
    if float(x) == int(x):
        return str(int(x))
    return repr(float(x))


def fmt_float(x):
    """Odoo-export style: always a float literal, e.g. 12.0."""
    return "" if x is None else repr(float(x))


def person_name(used):
    while True:
        style = rng.random()
        f, s = rng.choice(FIRST), rng.choice(SUR)
        if style < 0.45:
            n = f"{rng.choice(GIVEN)} {s} ({f})"
        elif style < 0.7:
            n = f"{rng.choice(GIVEN).upper()} {s.upper()}"
        else:
            n = f"{f} {s}"
        if n not in used:
            used.add(n)
            return n


# ---------------------------------------------------------------- products
def build_products():
    used_cn, used_codes, brands, products = set(), set(), [], []

    def new_brand(category, cn=None, en=None):
        while cn is None or cn in used_cn or en.upper() in AVOID_BRANDS:
            a, b = rng.sample(SYL, 2)
            cn, en = a[0] + b[0], (a[1] + b[1].lower())
        used_cn.add(cn)
        br = dict(cn=cn, en=en.upper() if rng.random() < 0.4 else en, category=category)
        brands.append(br)
        return br

    def code(prefix):
        while True:
            c = f"{prefix}{rng.randint(1, 99):02d}{rng.randint(101, 399)}"
            if c not in used_codes:
                used_codes.add(c)
                return c

    def pack(cat, unit, sub):
        if unit == "each":
            return f"{rng.choice([60, 100, 150, 200, 250, 400, 500])}ml"
        if unit == "Bag":
            return f"1bag*{rng.choice([5, 10, 20, 25, 50])}lb"
        if cat == "饮料":
            return f"{rng.choice([12, 15, 20, 24])}btls*{rng.choice([280, 330, 480, 500, 1500])}ml"
        if sub == "冰淇淋":
            return f"{rng.choice([8, 6, 24])}boxes*{rng.choice([3, 4])}bags*{rng.choice([72, 88, 90])}ml"
        return f"{rng.choice([6, 8, 10, 12, 20, 24, 30])}bags*{rng.choice([70, 100, 150, 200, 400, 500, 1000])}g"

    def make(br, sub_row, flavor, extra_cn=""):
        cat, sub, unit = sub_row["category"], sub_row["subcategory"], sub_row["unit"]
        noun = "Frozen Dessert" if sub == "冰淇淋" else CAT_NOUN.get(cat, "Food")
        name = f"{br['en']} {flavor[1]} {noun} | {br['cn']} {flavor[0]}{extra_cn}{sub} {pack(cat, unit, sub)}"
        lo, hi = PRICE_RANGE.get(cat, (15, 55))
        price = round(rng.uniform(lo, hi) * 2) / 2
        cost = round(price * rng.uniform(0.66, 0.9), 3)
        products.append(dict(ref=code(sub_row["prefix"]), name=name, brand=br, cat=cat, sub=sub, unit=unit,
                             shelf=sub_row["shelf_life"], price=price, cost=cost,
                             pop=rng.paretovariate(1.3), ice=(sub == "冰淇淋")))

    subs = {r["subcategory"]: r for r in TAX["categories"]}
    # stand-in brand for the single-brand report page (brand filter "清泉", product-name exclusion "茉莉")
    qq = new_brand("饮料", "清泉", "QINGQUAN")
    for i, (sub, fl) in enumerate([(s, f) for s in ("即饮茶", "植物饮料") for f in FLAVORS[5:12]]):
        make(qq, subs[sub], fl, "茉莉" if i in (1, 8) else "")
    for p in products:
        p["pop"] *= 3
    # ice cream brands
    for _ in range(3):
        br = new_brand("冷冻和冷藏食品")
        for fl in rng.sample(FLAVORS[3:12] + FLAVORS[25:34], rng.randint(8, 12)):
            make(br, subs["冰淇淋"], fl)
    for p in products:
        if p["ice"]:
            p["pop"] *= 4
    # everything else
    by_cat = defaultdict(list)
    for r in TAX["categories"]:
        if r["subcategory"] != "冰淇淋":
            by_cat[r["category"]].append(r)
    cat_weight = {c: sum(r["weight"] for r in rs) for c, rs in by_cat.items()}
    for _ in range(95):
        cat = rng.choices(list(cat_weight), list(cat_weight.values()))[0]
        br = new_brand(cat)
        rows = by_cat[cat]
        for _ in range(rng.randint(3, 9)):
            sub_row = rng.choices(rows, [r["weight"] for r in rows])[0]
            make(br, sub_row, rng.choice(FLAVORS))
    return brands, products


FEES = [
    dict(ref="Sales Discount", name="折扣"),
    dict(ref="C-1", name="C-1"),
    dict(ref="SHH", name="运费 Shipping & Handling"),
    dict(ref="Service Fee", name="Service Fee"),
    dict(ref="RDC", name="RDC至收货点运费"),
]


# ---------------------------------------------------------------- people & customers
def build_people_and_customers(products):
    used = set()
    reps, buyers, pickers = {}, {}, [person_name(used) for _ in range(5)]
    for co in COMPANIES:
        reps[co["name"]] = [person_name(used) for _ in range(co["reps"])]
        buyers[co["name"]] = [person_name(used) for _ in range(2)]
    reps["Crestline New Jersey Inc"].append("MAINSTREAM")
    leaders = {co["name"]: rng.sample(reps[co["name"]][:-1] if co["name"] == "Crestline New Jersey Inc" else reps[co["name"]], 1)[0]
               for co in COMPANIES}

    geo = defaultdict(list)
    for g in TAX["geo"]:
        geo[g["company"]].append(g)

    customers, used_names, seq = [], set(), 100
    for co in COMPANIES:
        places = geo[co["name"]]
        for i in range(co["customers"]):
            ctype = pick(CUST_TYPES)
            g = rng.choice(places)
            if co["name"] == "Crestline New Jersey Inc" and i < 14:
                g = next(p for p in places if p["city"] == "FLUSHING")
            while True:
                if ctype == "Employee":
                    nm = f"Employee {person_name(used)}"
                else:
                    nm = f"{rng.choice(STORE_A)} {rng.choice(STORE_B)} {rng.choice(TYPE_SUFFIX[ctype])}"
                if nm not in used_names:
                    used_names.add(nm)
                    break
            customers.append(new_customer(co, nm, ctype, g, reps, seq))
            seq += 1
        if co["name"] == "Crestline New Jersey Inc":
            nj = [p for p in places if p["state"].startswith("New Jersey")] or places
            for nm in ("Harbor Spice Co", "Summit Seasoning, LLC"):
                customers.append(new_customer(co, nm, "Wholesaler", rng.choice(nj), reps, seq))
                seq += 1
    for c in customers:
        pool = [p for p in products if not (c["type"] == "Pharmacy" and p["cat"] != "美容 & 健康 & 居家")]
        size = {"Wholesaler": 140, "Supermarket/Grocery": 110, "E-Commerce": 70}.get(c["type"], 45)
        size = max(8, int(size * rng.uniform(0.4, 1.4)))
        chosen = set()
        weights = [p["pop"] for p in pool]
        while len(chosen) < min(size, len(pool)):
            chosen.add(rng.choices(range(len(pool)), weights)[0])
        c["catalog"] = [pool[i] for i in chosen]
        c["cat_w"] = [p["pop"] for p in c["catalog"]]
    return reps, buyers, pickers, leaders, customers


def city_state_address(city, state):
    # Street and ZIP are drawn but not used, so the random sequence (and the rest of the data) stays the same.
    rng.randint(10, 9899), rng.choice(STREETS), rng.randint(1000, 9999)
    return f"{city}, {state}"


def new_customer(co, name, ctype, g, reps, seq):
    st = g["state"].replace(" (US)", "")
    ab = STATE_ABBR.get(st, "US")
    city = g["city"]
    team = reps[co["name"]]
    rep = "MAINSTREAM" if ("MAINSTREAM" in team and rng.random() < 0.06) else rng.choice([r for r in team if r != "MAINSTREAM"])
    start = START + timedelta(days=int(rng.random() ** 2.2 * 600)) if rng.random() < 0.35 else START
    if rng.random() < 0.05:
        start = date(2026, 8, 1) + timedelta(days=rng.randint(0, 45))
    end = AS_OF if rng.random() > 0.1 else START + timedelta(days=rng.randint(120, 560))
    freq = {"Wholesaler": 7, "Supermarket/Grocery": 5.5, "On-Premise": 3.5, "E-Commerce": 4, "Retail": 2.5,
            "Pharmacy": 2, "Others": 2, "Employee": 0.6}[ctype] * rng.lognormvariate(0, 0.55)
    return dict(
        company=co["name"], name=name, type=ctype, city=city, state=g["state"],
        no=f"US-{ab}{city.replace(' ', '')[:3].upper()}-{rng.randint(10, 9999)}",
        address=city_state_address(city, st),
        race=pick(RACES) if name not in ("Harbor Spice Co", "Summit Seasoning, LLC") else "", rep=rep,
        partner_rep=rep if rng.random() > 0.35 else "", start=start, end=end, freq=freq,
        big=ctype == "Wholesaler", disc=rng.uniform(0.08, 0.2) if ctype == "Wholesaler" else rng.uniform(0, 0.08),
    )


# ---------------------------------------------------------------- sales
def season(p, d):
    m = d.month
    if p["ice"]:
        return [0.25, 0.3, 0.55, 0.9, 1.6, 2.2, 2.4, 2.2, 1.3, 0.7, 0.35, 0.3][m - 1]
    if p["cat"] == "饮料":
        return [0.7, 0.75, 0.9, 1.0, 1.15, 1.3, 1.35, 1.3, 1.05, 0.9, 0.8, 0.85][m - 1]
    return [1.1, 1.25, 0.95, 0.95, 1.0, 0.95, 0.95, 1.0, 1.0, 1.05, 1.1, 1.2][m - 1]


def build_sales(customers, reps, pickers, leaders):
    lines, refunds = [], []
    seq = defaultdict(int)
    so_seq = defaultdict(int)
    day = START
    while day <= AS_OF:
        if day.weekday() == 6 and rng.random() > 0.05:
            day += timedelta(days=1)
            continue
        growth = 1.0 + 0.42 * max(0, (day - START).days) / 365
        for c in customers:
            if not (c["start"] <= day <= c["end"]):
                continue
            if rng.random() > c["freq"] / 26 * min(growth, 1.8) * 0.36:
                continue
            co = next(x for x in COMPANIES if x["name"] == c["company"])
            seq[co["inv"], day.year] += 1
            inv_no = f"{co['inv']}/{day.year}/{seq[co['inv'], day.year]:05d}"
            so_seq[co["so"], day.year] += 1
            so_ref = f"{co['so']}{day.strftime('%y%m')}{so_seq[co['so'], day.year] + 40000:05d}"
            creator = c["rep"] if rng.random() > 0.12 else rng.choice(reps[c["company"]])
            wh = EXTRA_WH.get(c["company"]) if (c["company"] in EXTRA_WH and rng.random() < 0.06) else co["wh"]
            n_lines = max(1, int(rng.gammavariate(2.2, 2.2 if c["big"] else 1.6)))
            chosen = set()
            for _ in range(n_lines * 3):
                if len(chosen) >= n_lines:
                    break
                p = rng.choices(c["catalog"], c["cat_w"])[0]
                if p["ref"] not in chosen and rng.random() < season(p, day) / 2.5 + 0.2:
                    chosen.add(p["ref"])
                    qty = max(1, int(rng.gammavariate(1.6, (22 if c["big"] else 4) * (1.4 if p["ice"] else 1))))
                    if p["unit"] == "each":
                        qty *= rng.choice([6, 12, 24])
                    price = round(p["price"] * (1 - c["disc"]) * rng.uniform(0.97, 1.03), 2)
                    cost = p["cost"] * (1.22 if (p["ice"] and day >= date(2025, 12, 1)) else 1.0)
                    if rng.random() < 0.01:
                        price = 0.0
                    lines.append(dict(inv=inv_no, date=day, p=p, c=c, qty=qty, price=price, cost=round(cost, 3),
                                      so=so_ref, creator=creator, wh=wh, kind="sale", co=co))
            r = rng.random()
            if r < 0.04:
                lines.append(dict(inv=inv_no, date=day, p=FEES[0], c=c, qty=1, price=round(rng.uniform(20, 400), 2),
                                  cost=0, so=so_ref, creator=creator, wh=wh, kind="discount", co=co))
            elif r < 0.06:
                lines.append(dict(inv=inv_no, date=day, p=FEES[2], c=c, qty=1, price=float(rng.choice([35, 75, 150, 200])),
                                  cost=0, so=so_ref, creator=creator, wh=wh, kind="fee", co=co))
            elif r < 0.063:
                lines.append(dict(inv=inv_no, date=day, p=FEES[1], c=c, qty=1, price=float(rng.choice([30, 60, 200])),
                                  cost=0, so=so_ref, creator=creator, wh=wh, kind="credit", co=co))
            elif r < 0.0635:
                lines.append(dict(inv=inv_no, date=day, p=FEES[3], c=c, qty=rng.choice([2, 6, 12]), price=450.0,
                                  cost=0, so=so_ref, creator=creator, wh=wh, kind="fee", co=co))
            if rng.random() < 0.018 and day < AS_OF - timedelta(days=3):
                src = [l for l in lines[-n_lines - 2:] if l["inv"] == inv_no and l["kind"] == "sale" and l["price"] > 0]
                if src:
                    rdate = day + timedelta(days=rng.randint(1, 10))
                    if rdate <= AS_OF:
                        refunds.append((rdate, src[0], max(1, int(src[0]["qty"] * rng.uniform(0.2, 1)))))
        day += timedelta(days=1)
    rseq = defaultdict(int)
    for rdate, l, q in sorted(refunds, key=lambda x: x[0]):
        co = l["co"]
        rseq[co["rinv"], rdate.year] += 1
        lines.append(dict(l, inv=f"{co['rinv']}/{rdate.year}/{rseq[co['rinv'], rdate.year]:05d}", date=rdate,
                          qty=q, kind="refund"))
    lines.sort(key=lambda l: (l["date"], l["inv"]), reverse=True)
    for l in lines:
        l["picker"] = rng.choice(pickers)
        l["leader"] = leaders[l["c"]["company"]] if rng.random() < 0.01 else ""
    return lines


def write_branch_sales(lines, path):
    header = ["lnvoice No", "Invoice Date", "Ref Num", "product_name", "in_charge", "brand_name", "brand_cn_name",
              "Category (CN)", "Sub-Category (CN)", "shelf life", "customer_name", "customer_no", "sales leader",
              "salesperson", "rep", "customer_type", "city", "state", "race", "quantity", "discount_price", "amount",
              "purchase_price", "cogs", "Gross Profit", "Gross Margin", "month", "week", "year", "company",
              "Warehouse", "Sales Order Line / Created By", "Sales Order Line / Order Reference",
              "Partner / Complete Address"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for l in lines:
            p, c, k = l["p"], l["c"], l["kind"]
            product = k == "sale" or k == "refund"
            sign = -1 if k in ("refund", "credit") else 1
            qty = sign * l["qty"]
            if k == "discount":
                qty, price, amount, cost = 1, 0.0, -l["price"], 0.0
            else:
                price = l["price"]
                amount = round(qty * price, 2) if price else None
                cost = l["cost"] if product else 0.0
            cogs = round(qty * cost, 3) if product else 0.0
            gp = (amount or 0) - cogs
            gm = "-∞" if not amount else repr(gp / amount)
            iso = l["date"].isocalendar()
            w.writerow([
                l["inv"], l["date"].isoformat(), p["ref"], p["name"], l["picker"],
                p["brand"]["en"] if product else "All", p["brand"]["cn"] if product else "All",
                p["cat"] if product else "All", p["sub"] if product else "All", p.get("shelf", "") if product else "",
                c["name"], c["no"], l["leader"], c["rep"], c["rep"], c["type"], c["city"], c["state"], c["race"],
                fmt_num(qty), fmt_num(price), fmt_num(amount) if amount is not None else "", fmt_num(cost),
                fmt_num(cogs), fmt_num(round(gp, 3)), gm, l["date"].month, iso[1], l["date"].year,
                c["company"], l["wh"], l["creator"], l["so"], c["address"],
            ])


def write_sales_all(lines, path):
    header = ["数量", "发票/账单日期", "产品/内部参考号", "产品/名称", "产品/CN 品牌", "产品/产品类别/中文名称",
              "产品/产品类别/上级名称", "产品/计量单位", "产品/基准成本", "合作伙伴", "合作伙伴/参考", "合作伙伴/销售员",
              "合作伙伴/客户类型", "合作伙伴/完整地址", "合作伙伴/城市", "合作伙伴/省/州", "销售订单行/订单参考号",
              "销售订单行/客户", "销售订单行/销售人员", "销售订单行/创建人", "销售订单行/仓库", "数量", "单价",
              "折扣后单价", "小计", "Type", "公司", "会计科目/类型"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
        w.writerow(header)
        for l in lines:
            if not l["co"]["in_sales_all"]:
                continue
            p, c, k = l["p"], l["c"], l["kind"]
            product = k in ("sale", "refund")
            ttype = "out_refund" if k in ("refund", "credit") else "out_invoice"
            qty = -1.0 if k == "discount" else float(l["qty"])
            price = float(l["price"])
            sub = "" if (price == 0 and k == "sale") else fmt_float(round(qty * price, 2))
            w.writerow([
                l["inv"], l["date"].isoformat(), p["ref"], p["name"],
                p["brand"]["cn"] if product else "", p["sub"] if product else "", p["cat"] if product else "",
                p["unit"] if product else "Units", fmt_float(p["cost"]) if product else "0.0",
                c["name"], c["no"], c["partner_rep"], c["type"], c["address"], c["city"], c["state"],
                l["so"], c["name"], c["rep"], l["creator"], l["wh"],
                fmt_float(qty), fmt_float(price), fmt_float(price), sub, ttype, c["company"], "收入",
            ])


# ---------------------------------------------------------------- inventory
def build_inventory(products, lines):
    demand = defaultdict(float)
    for l in lines:
        if l["kind"] == "sale" and l["date"] > AS_OF - timedelta(days=90):
            demand[l["co"]["name"], l["p"]["ref"]] += l["qty"] / 90
    lots, lot_seq = [], 0
    for co in COMPANIES:
        for p in products:
            d = demand.get((co["name"], p["ref"]), 0)
            if d == 0 and rng.random() > 0.12:
                continue
            state = rng.choices(["out", "low", "normal", "over", "dead"], [8, 15, 55, 17, 5])[0]
            cover = {"out": 0, "low": rng.uniform(4, 20), "normal": rng.uniform(30, 90),
                     "over": rng.uniform(120, 420), "dead": 0}[state]
            total = int(d * cover) if state != "dead" else rng.randint(5, 300)
            if d == 0:
                total = rng.randint(3, 250)
            n = rng.randint(1, 3)
            parts = [total // n] * n
            parts[0] += total - sum(parts)
            for q in parts:
                r = rng.random()
                if r < 0.05:
                    exp = AS_OF - timedelta(days=rng.randint(5, 150))
                elif r < 0.17:
                    exp = AS_OF + timedelta(days=rng.randint(3, 90))
                else:
                    exp = AS_OF + timedelta(days=rng.randint(91, max(120, int(p["shelf"] * 0.95))))
                lot_seq += 1
                lots.append(dict(co=co["name"], p=p, qty=max(0, q), exp=exp, lot=f"L{exp.strftime('%y%m')}{lot_seq:05d}"))
            for _ in range(rng.choice([0, 0, 1, 2])):
                lot_seq += 1
                exp = AS_OF - timedelta(days=rng.randint(30, 400))
                lots.append(dict(co=co["name"], p=p, qty=0, exp=exp, lot=exp.isoformat() if rng.random() < 0.5 else f"L{exp.strftime('%y%m')}{lot_seq:05d}"))
    return lots, demand


def write_lots(lots, transit, path):
    avail = defaultdict(int)
    for l in lots:
        avail[l["co"], l["p"]["ref"]] += l["qty"]
    pallet = {}
    header = ["批次/序列号码", "内部参考号", "产品/内部参考号", "产品/名称", "产品/产品类别/上级名称", "产品/产品类别/中文名称",
              "产品/CN 品牌", "显示名称", "在手数量", "有效期", "产品/批次可用", "产品/预测数量", "产品/打板",
              "产品/计量单位", "产品/基准成本", "活动", "产品/有效期", "产品/销售价格", "公司"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
        w.writerow(header)
        for l in lots:
            p = l["p"]
            if p["ref"] not in pallet:
                a, b = rng.randint(4, 18), rng.randint(5, 14)
                pallet[p["ref"]] = f"{a}*{b}={a * b}" if rng.random() < 0.7 else ""
            key = (l["co"], p["ref"])
            w.writerow([
                l["lot"], "", p["ref"], p["name"], p["cat"], p["sub"], p["brand"]["cn"],
                f"{l['exp'].strftime('%m-%d-%Y')}/0.0", fmt_float(l["qty"]), l["exp"].isoformat(),
                fmt_float(avail[key]), fmt_float(avail[key] + transit.get(key, 0)), pallet[p["ref"]], p["unit"],
                fmt_float(p["cost"]), "已到警报日期" if l["exp"] <= AS_OF + timedelta(days=60) and l["qty"] > 0 else "",
                str(p["shelf"]), fmt_float(999.0 if rng.random() < 0.03 else p["price"]), l["co"],
            ])


def write_branch_inventory(lots, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "库存明细"
    ws.append(["位置/仓库", "Brand", "产品/名称", "Ref Num", "批次/序列号码/有效期", "year_month", "数量1"])
    agg = defaultdict(int)
    for l in lots:
        if l["qty"] > 0:
            agg[l["co"], l["p"]["ref"], l["exp"]] += l["qty"]
    pmap = {l["p"]["ref"]: l["p"] for l in lots}
    for (co, ref, exp), q in sorted(agg.items(), key=lambda x: (x[0][0], x[0][1], x[0][2])):
        p = pmap[ref]
        ws.append([co, p["brand"]["en"], p["name"], ref, datetime(exp.year, exp.month, exp.day), exp.strftime("%Y-%m"), q])
    wb.save(path)


# ---------------------------------------------------------------- purchasing
def build_purchasing(brands, products, demand, buyers):
    suppliers, used = [], set()
    while len(suppliers) < 26:
        nm = rng.choice([
            f"{rng.choice(SUPPLIER_PLACES)} {rng.choice(SUPPLIER_WORDS)} Trading Co., LTD",
            f"{rng.choice(STORE_A)} {rng.choice(SUPPLIER_WORDS)} Imports Inc.",
            f"{rng.choice(SUPPLIER_PLACES)} Foods USA Inc",
        ])
        if nm in used:
            continue
        used.add(nm)
        imported = not nm.endswith("Inc.") or rng.random() < 0.3
        base = rng.uniform(35, 70) if imported else rng.uniform(3, 14)
        suppliers.append(dict(name=nm, lt=base, sd=base * rng.uniform(0.08, 0.3)))
    brand_sup = {b["cn"]: rng.choice(suppliers) for b in brands}
    by_sup = defaultdict(list)
    for p in products:
        by_sup[brand_sup[p["brand"]["cn"]]["name"]].append(p)
    sup_map = {s["name"]: s for s in suppliers}

    po_lines, supplements, transit = [], [], defaultdict(int)
    seq = defaultdict(int)
    for co in COMPANIES:
        for sname, prods in by_sup.items():
            s = sup_map[sname]
            stocked = [p for p in prods if demand.get((co["name"], p["ref"]), 0) > 0]
            if not stocked:
                continue
            cycle = rng.uniform(10, 35) if s["lt"] > 20 else rng.uniform(5, 16)
            day = datetime.combine(PO_START, datetime.min.time()) + timedelta(days=rng.uniform(0, cycle))
            while day.date() <= AS_OF:
                seq[co["po"]] += 1
                ref = f"{co['po']}{day.strftime('%y%m')}{seq[co['po']] + 8000:04d}"
                ref_full = ref + (f" (OL{rng.randint(1000, 4999)}-{rng.randint(1, 5)})" if rng.random() < 0.18 else "")
                confirm = datetime.combine(day.date(), datetime.min.time()) + timedelta(hours=rng.uniform(8, 20))
                expected = confirm + timedelta(days=round(s["lt"]))
                actual_lt = max(1.0, rng.gauss(s["lt"], s["sd"]))
                arrival = confirm + timedelta(days=actual_lt)
                arrived = arrival.date() <= AS_OF
                shipped = (confirm + timedelta(days=actual_lt * rng.uniform(0.1, 0.4))).date() if rng.random() < 0.12 else None
                if shipped and shipped > AS_OF:
                    shipped = None
                buyer = rng.choice(buyers[co["name"]])
                for p in rng.sample(stocked, min(len(stocked), rng.randint(1, 6))):
                    d = demand[co["name"], p["ref"]]
                    qty = max(1, int(d * cycle * rng.uniform(0.8, 1.5)))
                    if rng.random() < 0.01:
                        qty = -rng.randint(1, 5)
                    received = qty if arrived else 0
                    if arrived and qty > 5 and rng.random() < 0.04:
                        received = int(qty * rng.uniform(0.5, 0.95))
                    if arrived and rng.random() < 0.01:
                        received = 0
                    if not arrived and qty > 0:
                        transit[co["name"], p["ref"]] += qty
                    price = round(p["cost"] * rng.uniform(0.9, 1.02), 2)
                    po_lines.append([
                        ref_full, f"[{p['ref']}] {p['name']}", fmt_float(price), fmt_float(qty), fmt_float(received),
                        "" if rng.random() < 0.05 else fmt_float(round(qty * price, 2)),
                        confirm.strftime("%Y-%m-%d %H:%M:%S"), expected.strftime("%Y-%m-%d %H:%M:%S"), sname,
                        co["name"], buyer, "done" if arrived else "purchase",
                        shipped.isoformat() if (shipped and rng.random() < 0.5) else "",
                    ])
                supplements.append(dict(ref=ref, created=confirm - timedelta(minutes=rng.randint(1, 40)),
                                        shipped=shipped, arrival=arrival if arrived else None))
                day += timedelta(days=cycle * rng.uniform(0.7, 1.3))
    po_lines.sort(key=lambda r: r[6], reverse=True)
    supplements.sort(key=lambda r: r["created"])
    for i, s in enumerate(supplements):
        s["id"] = 20000 + i
    supplements.reverse()
    return po_lines, supplements, transit


def write_po(po_lines, supplements, po_path, sup_path):
    with open(po_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
        w.writerow(["订单行/订单关联", "订单行/产品", "订单行/单价", "订单行/数量", "订单行/已接收数量", "订单行/小计",
                    "订单行/确认日期", "订单行/预计到货时间", "订单行/合作伙伴", "订单行/公司", "订单行/创建人",
                    "订单行/状态", "订单行/订单关联/实际发货日期"])
        w.writerows(po_lines)
    with open(sup_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
        w.writerow(["ID", "订单关联", "创建日期", "实际发货日期", "到达"])
        for s in supplements:
            w.writerow([s["id"], s["ref"], s["created"].strftime("%Y-%m-%d %H:%M:%S"),
                        s["shipped"].isoformat() if s["shipped"] else "",
                        s["arrival"].strftime("%Y-%m-%d %H:%M:%S") if s["arrival"] else ""])


# ---------------------------------------------------------------- targets
def write_targets(lines, reps, path):
    nj = "Crestline New Jersey Inc"
    monthly = defaultdict(float)
    for l in lines:
        if l["c"]["company"] == nj and l["date"].year == 2026 and l["kind"] in ("sale", "refund"):
            sign = -1 if l["kind"] == "refund" else 1
            monthly[l["creator"], l["date"].month] += sign * l["qty"] * l["price"]
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI按月"
    ws.append(["月份日期", "YearMonth", "年份", "月份", "销售人员", "KPI目标", "公司"])
    raw = []
    people = [r for r in reps[nj] if r != "MAINSTREAM"]
    for r in people:
        avg = sum(monthly[r, m] for m in range(1, 9)) / 8
        base = max(20000, avg)
        targets = {}
        for m in range(1, 13):
            t = int(round(base * rng.uniform(0.9, 1.2) / 5000) * 5000)
            targets[m] = max(t, 10000)
            ws.append([datetime(2026, m, 1), f"2026-{m:02d}", 2026, m, r, targets[m], nj])
        raw.append([r, targets[8], targets[9]])
    ws2 = wb.create_sheet("原始数据")
    ws2.append(["Customer/Salesperson", "8月任务", "9月任务"])
    for row in raw:
        ws2.append(row)
    wb.save(path)


def ice_cream_target(lines):
    nj = [l for l in lines if l["c"]["company"] == "Crestline New Jersey Inc" and l["p"].get("ice") and l["kind"] == "sale"]
    by_m = defaultdict(float)
    for l in nj:
        if l["date"].year == 2026:
            by_m[l["date"].month] += l["qty"] * l["price"]
    peak = sum(by_m[m] for m in (6, 7, 8)) / 3
    return int(round(peak * 1.1 / 10000) * 10000)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    brands, products = build_products()
    reps, buyers, pickers, leaders, customers = build_people_and_customers(products)
    lines = build_sales(customers, reps, pickers, leaders)
    write_branch_sales(lines, OUT / "branch_sales.csv")
    write_sales_all(lines, OUT / "sales_invoice_lines.csv")
    lots, demand = build_inventory(products, lines)
    po_lines, supplements, transit = build_purchasing(brands, products, demand, buyers)
    write_lots(lots, transit, OUT / "inventory_lots.csv")
    write_branch_inventory(lots, OUT / "branch_inventory.xlsx")
    write_po(po_lines, supplements, OUT / "purchase_order_lines.csv", OUT / "purchase_tracking.csv")
    write_targets(lines, reps, OUT / "sales_targets_monthly.xlsx")
    print(f"brands {len(brands)}, products {len(products)}, customers {len(customers)}, sales lines {len(lines)}, "
          f"lots {len(lots)}, PO lines {len(po_lines)}, POs {len(supplements)}")
    print(f"suggested 冰淇淋总目标 = {ice_cream_target(lines)}")


if __name__ == "__main__":
    main()
